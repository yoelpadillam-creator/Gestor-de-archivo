"""
Exporta el informe a Excel (openpyxl) o PDF (reportlab).
"""

import io
from datetime import datetime

from fastapi import APIRouter
from fastapi.responses import StreamingResponse

router = APIRouter(prefix="/api/export")


# ── Excel ──────────────────────────────────────────────────────────────────

def _build_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment

    wb = Workbook()
    wb.remove(wb.active)   # eliminar hoja por defecto

    HDR_FILL = PatternFill("solid", fgColor="2D6A4F")
    HDR_FONT = Font(color="FFFFFF", bold=True)
    HDR_ALIGN = Alignment(horizontal="center")

    def add_sheet(name: str, headers: list, rows: list):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = HDR_FILL
            cell.font = HDR_FONT
            cell.alignment = HDR_ALIGN
        for row in rows:
            ws.append([str(v) if v is not None else "" for v in row])

    s = report['summary']
    add_sheet("Resumen", ["Métrica", "Valor"], [
        ["Facturas analizadas",          s['invoices']],
        ["Entradas detectadas",          s['entries_detected']],
        ["Filtro de modelo",             s['model_filter']],
        ["Facturas duplicadas",          s['duplicates']],
        ["Coincidencias exactas",        s['matches_exact']],
        ["Coincidencias parciales",      s['matches_partial']],
        ["Facturados sin inventario",    s['missing_in_inventory']],
        ["En inventario no facturados",  s['not_invoiced']],
        ["Discrepancias color (FIT)",    s['fit_color_discrepancies']],
    ])

    dup_rows = []
    for d in report['duplicates']:
        for e in d['entries']:
            dup_rows.append([d['by'], d['key'], e['file'], e['page'],
                              e.get('modelo'), e.get('chasis'), e.get('motor'),
                              e.get('color')])
    add_sheet("A_Duplicados",
              ["Campo", "Valor", "Factura", "Pág.", "Modelo", "Chasis", "Motor", "Color"],
              dup_rows)

    add_sheet("B_Coincidencias",
              ["Factura", "Pág.", "Chasis Fac.", "Chasis Inv.", "Motor Fac.",
               "Motor Inv.", "Color Fac.", "Color Inv.", "Chasis OK", "Motor OK",
               "Color OK", "Estado"],
              [[m['file'], m['page'],
                m.get('invoice_chasis'),  m.get('inventory_chasis'),
                m.get('invoice_motor'),   m.get('inventory_motor'),
                m.get('invoice_color'),   m.get('inventory_color'),
                "Sí" if m['chasis_ok'] else "No",
                "Sí" if m['motor_ok']  else "No",
                "Sí" if m['color_ok']  else "No",
                m['status']] for m in report['matches']])

    add_sheet("C_SinInventario",
              ["Factura", "Pág.", "Modelo", "Chasis", "Motor", "Color"],
              [[e['file'], e['page'], e.get('modelo'), e.get('chasis'),
                e.get('motor'), e.get('color')]
               for e in report['missing_in_inventory']])

    add_sheet("D_NoFacturados",
              ["Modelo", "Chasis", "Motor", "Color"],
              [[r.get('modelo'), r.get('chasis'), r.get('motor'), r.get('color')]
               for r in report['not_invoiced']])

    add_sheet("E_FIT_Colores",
              ["Chasis", "Color Factura", "Color FIT", "Modelo FIT", "Factura", "Pág."],
              [[d['chasis'], d.get('color_factura'), d.get('color_fit'),
                d.get('fit_modelo'), d['file'], d['page']]
               for d in report['fit_discrepancies']])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF ────────────────────────────────────────────────────────────────────

def _build_pdf(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph,
        Spacer, PageBreak
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    GREEN  = colors.HexColor("#2D6A4F")
    BEIGE  = colors.HexColor("#F0EDE8")
    BLACK  = colors.HexColor("#1A1A1A")

    def h1(text): return Paragraph(f"<b>{text}</b>", styles['Heading1'])
    def h2(text): return Paragraph(f"<b>{text}</b>", styles['Heading2'])
    def body(text): return Paragraph(text, styles['Normal'])

    def make_table(headers: list, rows: list) -> Table:
        data = [headers] + [[str(v) if v is not None else "—" for v in r]
                             for r in rows]
        t = Table(data, repeatRows=1, hAlign='LEFT')
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), GREEN),
            ('TEXTCOLOR',  (0, 0), (-1, 0), colors.white),
            ('FONTNAME',   (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE',   (0, 0), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BEIGE]),
            ('TEXTCOLOR',  (0, 1), (-1, -1), BLACK),
            ('GRID',       (0, 0), (-1, -1), 0.3, colors.HexColor("#D9D3CC")),
            ('VALIGN',     (0, 0), (-1, -1), 'TOP'),
            ('PADDING',    (0, 0), (-1, -1), 4),
        ]))
        return t

    story = []
    s = report['summary']
    story.append(h1("Informe de Facturas vs Inventario"))
    story.append(body(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}  |  "
                      f"Modelo: {s['model_filter']}"))
    story.append(Spacer(1, 0.4*cm))

    # Resumen
    story.append(make_table(
        ["Métrica", "Valor"],
        [["Facturas analizadas", s['invoices']],
         ["Entradas detectadas", s['entries_detected']],
         ["Duplicados", s['duplicates']],
         ["Coincidencias exactas", s['matches_exact']],
         ["Coincidencias parciales", s['matches_partial']],
         ["Facturados sin inventario", s['missing_in_inventory']],
         ["En inventario no facturados", s['not_invoiced']],
         ["Discrepancias color FIT", s['fit_color_discrepancies']]]
    ))

    sections = [
        ("A · Facturas duplicadas",
         ["Campo", "Valor", "Factura", "Pág.", "Chasis", "Motor"],
         [[d['by'], d['key'], e['file'], e['page'],
           e.get('chasis'), e.get('motor')]
          for d in report['duplicates'] for e in d['entries']]),

        ("B · Coincidencias con inventario",
         ["Factura", "Pág.", "Chasis Fac.", "Chasis Inv.", "Color Fac.", "Color Inv.", "Estado"],
         [[m['file'], m['page'], m.get('invoice_chasis'), m.get('inventory_chasis'),
           m.get('invoice_color'), m.get('inventory_color'), m['status']]
          for m in report['matches']]),

        ("C · Facturados sin inventario",
         ["Factura", "Pág.", "Modelo", "Chasis", "Motor", "Color"],
         [[e['file'], e['page'], e.get('modelo'), e.get('chasis'),
           e.get('motor'), e.get('color')]
          for e in report['missing_in_inventory']]),

        ("D · En inventario no facturados",
         ["Modelo", "Chasis", "Motor", "Color"],
         [[r.get('modelo'), r.get('chasis'), r.get('motor'), r.get('color')]
          for r in report['not_invoiced']]),

        ("E · Discrepancias de color (FIT)",
         ["Chasis", "Color Factura", "Color FIT", "Modelo FIT", "Factura", "Pág."],
         [[d['chasis'], d.get('color_factura'), d.get('color_fit'),
           d.get('fit_modelo'), d['file'], d['page']]
          for d in report['fit_discrepancies']]),
    ]

    for title, headers, rows in sections:
        story.append(PageBreak())
        story.append(h2(title))
        story.append(Spacer(1, 0.3*cm))
        if rows:
            story.append(make_table(headers, rows))
        else:
            story.append(body("<i>Sin datos.</i>"))

    doc.build(story)
    return buf.getvalue()


# ── Endpoints ──────────────────────────────────────────────────────────────

@router.post("/xlsx")
async def export_xlsx(report: dict):
    data = _build_xlsx(report)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=informe-facturas.xlsx"}
    )


@router.post("/pdf")
async def export_pdf(report: dict):
    data = _build_pdf(report)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=informe-facturas.pdf"}
    )
